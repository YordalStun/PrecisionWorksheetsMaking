import os
import tempfile
import unittest

from docx import Document
from docx.shared import Cm

from precision_worksheets.large_print import (
    build_large_print_pages,
    export_large_print_docx,
    export_large_print_pdf,
    render_large_print_image,
)


class TestBuildLargePrintPages(unittest.TestCase):
    def test_one_page_per_word_in_order(self):
        pages = build_large_print_pages(["cat", "dog", "sun"])
        self.assertEqual([p.word for p in pages], ["cat", "dog", "sun"])

    def test_blank_words_are_dropped(self):
        pages = build_large_print_pages(["cat", "  ", "sun"])
        self.assertEqual([p.word for p in pages], ["cat", "sun"])

    def test_rejects_empty_word_list(self):
        with self.assertRaises(ValueError):
            build_large_print_pages([])
        with self.assertRaises(ValueError):
            build_large_print_pages(["   "])


class TestRenderLargePrintImage(unittest.TestCase):
    def test_image_fills_a_landscape_shaped_canvas(self):
        img = render_large_print_image("cat")
        self.assertGreater(img.width, img.height)

    def test_word_ink_is_roughly_centred(self):
        img = render_large_print_image("jump")
        gray = img.convert("L")
        pixels = gray.load()
        dark_xs = [x for x in range(img.width) for y in range(img.height) if pixels[x, y] < 128]
        dark_ys = [y for x in range(img.width) for y in range(img.height) if pixels[x, y] < 128]
        ink_center_x = (min(dark_xs) + max(dark_xs)) / 2
        ink_center_y = (min(dark_ys) + max(dark_ys)) / 2
        self.assertAlmostEqual(ink_center_x, img.width / 2, delta=img.width * 0.03)
        self.assertAlmostEqual(ink_center_y, img.height / 2, delta=img.height * 0.03)

    def test_long_word_still_fits_within_the_canvas(self):
        img = render_large_print_image("extraordinarily")
        gray = img.convert("L")
        pixels = gray.load()
        dark_xs = [x for x in range(img.width) for y in range(img.height) if pixels[x, y] < 128]
        self.assertGreater(min(dark_xs), 0)
        self.assertLess(max(dark_xs), img.width)


class TestLargePrintExport(unittest.TestCase):
    def setUp(self):
        self.pages = build_large_print_pages(["cat", "jump", "sun"])
        self.tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmpdir.cleanup)

    def test_export_pdf_has_one_landscape_page_per_word(self):
        path = os.path.join(self.tmpdir.name, "large_print.pdf")
        export_large_print_pdf(self.pages, path)
        with open(path, "rb") as f:
            data = f.read()
        self.assertTrue(data.startswith(b"%PDF-"))
        self.assertEqual(data.count(b"/Type /Page"), len(self.pages) + 1)

    def test_export_pdf_rejects_empty_list(self):
        path = os.path.join(self.tmpdir.name, "large_print.pdf")
        with self.assertRaises(ValueError):
            export_large_print_pdf([], path)

    def test_export_docx_is_landscape_a4_with_one_paragraph_per_word(self):
        path = os.path.join(self.tmpdir.name, "large_print.docx")
        export_large_print_docx(self.pages, path)
        doc = Document(path)
        section = doc.sections[0]
        self.assertAlmostEqual(section.page_width / Cm(1), 29.7, places=1)
        self.assertAlmostEqual(section.page_height / Cm(1), 21.0, places=1)
        self.assertEqual(len(doc.paragraphs), len(self.pages))
        for paragraph in doc.paragraphs:
            self.assertEqual(len(paragraph.runs[0].element.findall(
                ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"
            )), 1)

    def test_export_docx_rejects_empty_list(self):
        path = os.path.join(self.tmpdir.name, "large_print.docx")
        with self.assertRaises(ValueError):
            export_large_print_docx([], path)


if __name__ == "__main__":
    unittest.main()
