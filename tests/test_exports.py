import os
import tempfile
import unittest

from docx import Document

from precision_worksheets.docx_export import export_docx
from precision_worksheets.generator import build_probe_sheets
from precision_worksheets.pdf_export import export_pdf


class TestExports(unittest.TestCase):
    def setUp(self):
        self.sheets = build_probe_sheets(
            "Alex", ["cat", "dog", "sun", "run", "big"], num_sheets=3,
            rows=10, cols=8, date_str="04/09/2026", seed=42,
        )
        self.tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmpdir.cleanup)

    def test_export_pdf_creates_valid_multi_page_file(self):
        path = os.path.join(self.tmpdir.name, "out.pdf")
        export_pdf(self.sheets, path)
        with open(path, "rb") as f:
            header = f.read(5)
            self.assertEqual(header, b"%PDF-")
        self.assertGreater(os.path.getsize(path), 0)

    def test_export_pdf_rejects_empty_sheet_list(self):
        path = os.path.join(self.tmpdir.name, "out.pdf")
        with self.assertRaises(ValueError):
            export_pdf([], path)

    def test_export_docx_creates_one_grid_and_one_tracker_table_per_sheet(self):
        path = os.path.join(self.tmpdir.name, "out.docx")
        export_docx(self.sheets, path)
        doc = Document(path)
        # A grid table and a progress-tracker table for each of the 3 sheets.
        self.assertEqual(len(doc.tables), 6)
        grid_tables = doc.tables[0::2]
        tracker_tables = doc.tables[1::2]
        for table in grid_tables:
            self.assertEqual(len(table.rows), 10)
            self.assertEqual(len(table.columns), 9)  # 8 words + row-number column
        for table in tracker_tables:
            self.assertEqual(len(table.rows), 3)  # Try, Date, Score
            self.assertEqual(len(table.columns), 11)  # label + 10 tries

    def test_export_docx_word_appears_in_grid(self):
        path = os.path.join(self.tmpdir.name, "out.docx")
        export_docx(self.sheets, path)
        doc = Document(path)
        cell_text = doc.tables[0].rows[0].cells[1].text
        self.assertIn(cell_text, ["cat", "dog", "sun", "run", "big"])

    def test_export_docx_rejects_empty_sheet_list(self):
        path = os.path.join(self.tmpdir.name, "out.docx")
        with self.assertRaises(ValueError):
            export_docx([], path)

    def test_export_pdf_stays_one_page_per_sheet_with_tracker_included(self):
        path = os.path.join(self.tmpdir.name, "out.pdf")
        export_pdf(self.sheets, path)
        with open(path, "rb") as f:
            data = f.read()
        # One page per sheet (tracker lives on the same page as its grid),
        # plus the "/Type /Pages" tree node itself also matching this
        # substring.
        self.assertEqual(data.count(b"/Type /Page"), len(self.sheets) + 1)

    def test_export_pdf_without_tracker_is_still_one_page_per_sheet(self):
        path = os.path.join(self.tmpdir.name, "out.pdf")
        export_pdf(self.sheets, path, include_tracker=False)
        with open(path, "rb") as f:
            data = f.read()
        self.assertEqual(data.count(b"/Type /Page"), len(self.sheets) + 1)

    def test_export_docx_tracker_table_is_try_date_score_transposed(self):
        path = os.path.join(self.tmpdir.name, "out.docx")
        export_docx(self.sheets, path)
        doc = Document(path)
        tracker = doc.tables[1]  # grid, tracker, grid, tracker, ...
        self.assertEqual(len(tracker.rows), 3)
        self.assertEqual(len(tracker.columns), 11)

        try_row, date_row, score_row = tracker.rows
        self.assertEqual(try_row.cells[0].text, "Try")
        self.assertEqual([c.text for c in try_row.cells[1:]], [str(i) for i in range(1, 11)])
        self.assertEqual(date_row.cells[0].text, "Date")
        self.assertEqual(score_row.cells[0].text, "Score")
        for cell in date_row.cells[1:]:
            self.assertEqual(cell.text, "")
        for cell in score_row.cells[1:]:
            self.assertEqual(cell.text, "")

    def test_export_docx_without_tracker_has_only_grid_tables(self):
        path = os.path.join(self.tmpdir.name, "out.docx")
        export_docx(self.sheets, path, include_tracker=False)
        doc = Document(path)
        self.assertEqual(len(doc.tables), len(self.sheets))
        for table in doc.tables:
            self.assertEqual(len(table.columns), 9)


if __name__ == "__main__":
    unittest.main()
