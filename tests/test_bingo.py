import os
import tempfile
import unittest
from collections import Counter

from docx import Document

from precision_worksheets.games.bingo import CARD_COLS, CARD_ROWS, build_bingo_cards, export_bingo_docx, export_bingo_pdf


class TestBuildBingoCards(unittest.TestCase):
    def test_builds_requested_number_of_cards(self):
        cards = build_bingo_cards("Alex", ["cat", "dog", "sun", "run", "big"], "04/09/2026", num_cards=4, seed=1)
        self.assertEqual(len(cards), 4)
        for i, card in enumerate(cards, start=1):
            self.assertEqual(card.card_number, i)
            self.assertEqual(card.total_cards, 4)

    def test_grid_is_card_rows_by_card_cols(self):
        cards = build_bingo_cards("Alex", ["cat", "dog", "sun", "run", "big"], "04/09/2026", num_cards=1, seed=1)
        grid = cards[0].grid
        self.assertEqual(len(grid), CARD_ROWS)
        self.assertTrue(all(len(row) == CARD_COLS for row in grid))

    def test_every_word_appears_and_no_adjacent_repeats(self):
        words = ["cat", "dog", "sun", "run", "big"]
        cards = build_bingo_cards("Alex", words, "04/09/2026", num_cards=1, seed=1)
        flat = [w for row in cards[0].grid for w in row]
        self.assertEqual(set(flat), set(words))
        for i in range(1, len(flat)):
            self.assertNotEqual(flat[i], flat[i - 1])

    def test_cards_are_independently_shuffled(self):
        cards = build_bingo_cards("Alex", ["a", "b", "c", "d", "e"], "04/09/2026", num_cards=4, seed=1)
        flattened = {tuple(w for row in c.grid for w in row) for c in cards}
        self.assertEqual(len(flattened), 4)

    def test_rejects_empty_word_list(self):
        with self.assertRaises(ValueError):
            build_bingo_cards("Alex", [], "04/09/2026")

    def test_rejects_zero_cards(self):
        with self.assertRaises(ValueError):
            build_bingo_cards("Alex", ["a"], "04/09/2026", num_cards=0)


class TestBingoExport(unittest.TestCase):
    def setUp(self):
        self.cards = build_bingo_cards("Alex", ["cat", "dog", "sun", "run", "big"], "04/09/2026", num_cards=3, seed=1)
        self.tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmpdir.cleanup)

    def test_export_pdf_has_one_page_per_card(self):
        path = os.path.join(self.tmpdir.name, "bingo.pdf")
        export_bingo_pdf(self.cards, path)
        with open(path, "rb") as f:
            data = f.read()
        self.assertTrue(data.startswith(b"%PDF-"))
        self.assertEqual(data.count(b"/Type /Page"), len(self.cards) + 1)

    def test_export_pdf_rejects_empty_list(self):
        path = os.path.join(self.tmpdir.name, "bingo.pdf")
        with self.assertRaises(ValueError):
            export_bingo_pdf([], path)

    def test_export_docx_has_one_table_per_card(self):
        path = os.path.join(self.tmpdir.name, "bingo.docx")
        export_bingo_docx(self.cards, path)
        doc = Document(path)
        self.assertEqual(len(doc.tables), len(self.cards))
        for table, card in zip(doc.tables, self.cards):
            self.assertEqual(len(table.rows), CARD_ROWS)
            self.assertEqual(len(table.columns), CARD_COLS)
            cell_words = [table.rows[r].cells[c].text for r in range(CARD_ROWS) for c in range(CARD_COLS)]
            self.assertEqual(Counter(cell_words), Counter(w for row in card.grid for w in row))


if __name__ == "__main__":
    unittest.main()
