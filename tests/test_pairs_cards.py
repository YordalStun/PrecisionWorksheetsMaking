import os
import tempfile
import unittest
from collections import Counter

from docx import Document

from precision_worksheets.games.pairs_cards import (
    CARD_COLS,
    CARDS_PER_WORD,
    build_pairs_card_sheet,
    export_pairs_cards_docx,
    export_pairs_cards_pdf,
)


class TestBuildPairsCardSheet(unittest.TestCase):
    def test_each_word_appears_twice(self):
        sheet = build_pairs_card_sheet(
            "Alex", ["cat", "dog", "sun", "run", "big"], "04/09/2026", seed=1
        )
        flat = [w for row in sheet.cards for w in row]
        counts = Counter(flat)
        self.assertEqual(set(counts), {"cat", "dog", "sun", "run", "big"})
        self.assertTrue(all(c == CARDS_PER_WORD for c in counts.values()))

    def test_grid_is_two_columns_wide(self):
        sheet = build_pairs_card_sheet("Alex", ["a", "b", "c"], "04/09/2026", seed=1)
        self.assertTrue(all(len(row) == CARD_COLS for row in sheet.cards))

    def test_no_adjacent_matching_cards(self):
        for seed in range(10):
            sheet = build_pairs_card_sheet(
                "Alex", ["cat", "dog", "sun", "run", "big"], "04/09/2026", seed=seed
            )
            flat = [w for row in sheet.cards for w in row]
            for i in range(1, len(flat)):
                self.assertNotEqual(flat[i], flat[i - 1], f"seed={seed} pos={i}")

    def test_rejects_empty_word_list(self):
        with self.assertRaises(ValueError):
            build_pairs_card_sheet("Alex", [], "04/09/2026")


class TestPairsCardsExport(unittest.TestCase):
    def setUp(self):
        self.sheet = build_pairs_card_sheet(
            "Alex", ["cat", "dog", "sun", "run", "big"], "04/09/2026", seed=1
        )
        self.tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmpdir.cleanup)

    def test_export_pdf_is_a_single_page(self):
        path = os.path.join(self.tmpdir.name, "cards.pdf")
        export_pairs_cards_pdf(self.sheet, path)
        with open(path, "rb") as f:
            data = f.read()
        self.assertTrue(data.startswith(b"%PDF-"))
        # "/Type /Page" also matches as a substring of "/Type /Pages" (the
        # page-tree root), so one page in the document means two hits here.
        self.assertEqual(data.count(b"/Type /Page"), 2)

    def test_export_docx_has_one_table_with_all_cards(self):
        path = os.path.join(self.tmpdir.name, "cards.docx")
        export_pairs_cards_docx(self.sheet, path)
        doc = Document(path)
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), self.sheet.rows)
        self.assertEqual(len(table.columns), CARD_COLS)
        cell_words = [table.rows[r].cells[c].text for r in range(len(table.rows)) for c in range(CARD_COLS)]
        self.assertEqual(Counter(cell_words), Counter(w for row in self.sheet.cards for w in row))


if __name__ == "__main__":
    unittest.main()
