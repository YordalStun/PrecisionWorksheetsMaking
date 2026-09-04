"""Bingo using the target words: each player gets a card with the 5 words
repeated in a shuffled grid. A caller reads words out loud (from the same
5-word list); mark every matching cell each time. First full line - across,
down, or diagonal - wins.
"""

from __future__ import annotations

import random
from dataclasses import dataclass

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from reportlab.lib.units import cm as POINTS_PER_CM

from ..docx_helpers import set_cell_margins, set_cell_text, set_column_widths, set_run_font
from ..fonts import get_comic_font_names
from ..generator import generate_word_grid
from ..layout import fit_grid_font_size

CARD_ROWS = 4
CARD_COLS = 4
CARD_FONT_CAP = 60

PAGE_WIDTH, PAGE_HEIGHT = A4
MARGIN = 12 * mm

INSTRUCTIONS = (
    "Cross off every matching word each time it's called. First to complete "
    "a full line - across, down, or diagonal - shouts BINGO!"
)


@dataclass
class BingoCard:
    child_name: str
    date_str: str
    words: list[str]
    grid: list[list[str]]
    card_number: int
    total_cards: int


def build_bingo_cards(
    child_name: str,
    words: list[str],
    date_str: str,
    num_cards: int = 4,
    seed: int | None = None,
) -> list[BingoCard]:
    """Build `num_cards` independently-shuffled bingo cards."""
    clean_words = [w.strip() for w in words if w.strip()]
    if not clean_words:
        raise ValueError("At least one word is required")
    if num_cards < 1:
        raise ValueError("num_cards must be at least 1")

    rng = random.Random(seed)
    cards = []
    for n in range(1, num_cards + 1):
        grid = generate_word_grid(clean_words, CARD_ROWS, CARD_COLS, rng)
        cards.append(
            BingoCard(
                child_name=child_name, date_str=date_str, words=clean_words,
                grid=grid, card_number=n, total_cards=num_cards,
            )
        )
    return cards


def export_bingo_pdf(cards: list[BingoCard], output_path: str) -> None:
    if not cards:
        raise ValueError("No cards to export")

    doc = SimpleDocTemplate(
        output_path, pagesize=A4,
        topMargin=MARGIN, bottomMargin=MARGIN, leftMargin=MARGIN, rightMargin=MARGIN,
        title="Bingo",
    )
    comic_regular, comic_bold = get_comic_font_names()
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "BGTitle", parent=styles["Title"], fontName=comic_bold, fontSize=26,
        alignment=TA_CENTER, spaceAfter=4,
    )
    meta_style = ParagraphStyle(
        "BGMeta", parent=styles["Normal"], fontName=comic_regular, fontSize=11,
        alignment=TA_CENTER, spaceAfter=6,
    )
    instructions_style = ParagraphStyle(
        "BGInstructions", parent=styles["Normal"], fontName=comic_regular, fontSize=10,
        alignment=TA_CENTER, textColor=colors.HexColor("#444444"), spaceAfter=10,
    )

    usable_width = PAGE_WIDTH - 2 * MARGIN

    story = []
    for card in cards:
        header_flowables = [
            Paragraph("BINGO!", title_style),
            Paragraph(
                f"Card {card.card_number} of {card.total_cards} &nbsp;&nbsp;&nbsp; "
                f"Date: {card.date_str}", meta_style,
            ),
            Paragraph(INSTRUCTIONS, instructions_style),
            Spacer(1, 4),
        ]
        story.extend(header_flowables)
        header_height = sum(f.wrap(usable_width, PAGE_HEIGHT)[1] for f in header_flowables)

        col_width = usable_width / CARD_COLS
        fitted_font_size = fit_grid_font_size(card.words, CARD_FONT_CAP, col_width)
        cell_style = ParagraphStyle(
            "BGCell", parent=styles["Normal"], fontName=comic_bold,
            fontSize=fitted_font_size, leading=fitted_font_size * 1.15, alignment=TA_CENTER,
        )

        available_height = (PAGE_HEIGHT - 2 * MARGIN) - header_height - 60
        target_row_height = available_height / CARD_ROWS
        text_block_height = fitted_font_size * 1.15
        vertical_padding = max(10, (target_row_height - text_block_height) / 2)

        data = [[Paragraph(word, cell_style) for word in row] for row in card.grid]
        table = Table(data, colWidths=[col_width] * CARD_COLS)
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 1.5, colors.HexColor("#333333")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), vertical_padding),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), vertical_padding),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(table)

        if card.card_number != card.total_cards:
            story.append(PageBreak())

    doc.build(story)


def export_bingo_docx(cards: list[BingoCard], output_path: str) -> None:
    if not cards:
        raise ValueError("No cards to export")

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    usable_width_cm = (section.page_width - section.left_margin - section.right_margin) / Cm(1)

    for card in cards:
        title = doc.add_heading("BINGO!", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            set_run_font(run)

        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta.add_run(f"Card {card.card_number} of {card.total_cards}     Date: {card.date_str}")
        meta_run.bold = True
        set_run_font(meta_run)

        instructions = doc.add_paragraph()
        instructions.alignment = WD_ALIGN_PARAGRAPH.CENTER
        instructions_run = instructions.add_run(INSTRUCTIONS)
        instructions_run.italic = True
        set_run_font(instructions_run)

        table = doc.add_table(rows=CARD_ROWS, cols=CARD_COLS)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.autofit = False

        col_width_cm = usable_width_cm / CARD_COLS
        set_column_widths(table, [col_width_cm] * CARD_COLS)

        fitted_font_size = fit_grid_font_size(card.words, CARD_FONT_CAP, col_width_cm * POINTS_PER_CM)

        header_reserve_pt = 200
        page_height_pt = 29.7 * POINTS_PER_CM
        margins_pt = (section.top_margin + section.bottom_margin) / Cm(1) * POINTS_PER_CM
        available_height = page_height_pt - margins_pt - header_reserve_pt
        target_row_height = available_height / CARD_ROWS
        text_block_height = fitted_font_size * 1.15
        vertical_padding = max(10, (target_row_height - text_block_height) / 2)
        set_cell_margins(table, top_pt=vertical_padding, bottom_pt=vertical_padding, left_pt=6, right_pt=6)

        for r, row_words in enumerate(card.grid):
            cells = table.rows[r].cells
            for c, word in enumerate(row_words):
                set_cell_text(cells[c], word, size_pt=fitted_font_size, bold=True)

        if card.card_number != card.total_cards:
            doc.add_page_break()

    doc.save(output_path)
