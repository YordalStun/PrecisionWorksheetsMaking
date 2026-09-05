"""Shared Pillow/export helpers for games that draw a whole page (board,
trail, ...) as one image with Pillow and then embed it, unchanged, into
both a PDF and a Word doc - so the two formats always look identical.
"""

from __future__ import annotations

import io

from PIL import Image, ImageDraw, ImageFont


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: float) -> list[str]:
    """Greedily wrap `text` onto lines that each fit within max_width."""
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        box = draw.textbbox((0, 0), candidate, font=font)
        if box[2] - box[0] <= max_width or not current:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def fit_font(
    draw: ImageDraw.ImageDraw, text: str, font_path: str, max_size: int, max_width: int, min_size: int = 14
) -> ImageFont.FreeTypeFont:
    """The largest font at `font_path`, up to max_size, at which `text`
    still fits within max_width."""
    size = max_size
    while size > min_size:
        font = ImageFont.truetype(font_path, size)
        box = draw.textbbox((0, 0), text, font=font)
        if box[2] - box[0] <= max_width:
            return font
        size -= 2
    return ImageFont.truetype(font_path, min_size)


def image_to_png_bytes(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


_DOCX_PAGE_SIZES_CM = {"A4": (21.0, 29.7), "A3": (29.7, 42.0)}

# SimpleDocTemplate's default page template wraps content in a Frame with
# 6pt of its own padding on every side (on top of the margins we set), so
# the space actually available for a flowable is 12pt narrower/shorter than
# (page size - 2*margin) in each dimension.
FRAME_PADDING = 12


def export_full_page_image_pdf(
    img: Image.Image, output_path: str, title: str, margin_mm: float = 8, page_size: str = "A4"
) -> None:
    """Embed `img` as a single, page-filling image in a new PDF."""
    from reportlab.lib.pagesizes import A3, A4
    from reportlab.lib.units import mm
    from reportlab.platypus import Image as RLImage, SimpleDocTemplate

    pagesize = {"A4": A4, "A3": A3}[page_size]
    png_bytes = image_to_png_bytes(img)
    margin = margin_mm * mm
    doc = SimpleDocTemplate(
        output_path, pagesize=pagesize,
        topMargin=margin, bottomMargin=margin, leftMargin=margin, rightMargin=margin,
        title=title,
    )
    page_width, page_height = pagesize
    usable_width = page_width - 2 * margin - FRAME_PADDING
    usable_height = page_height - 2 * margin - FRAME_PADDING
    aspect = img.height / img.width
    draw_width = usable_width
    draw_height = draw_width * aspect
    if draw_height > usable_height:
        draw_height = usable_height
        draw_width = draw_height / aspect

    rl_image = RLImage(io.BytesIO(png_bytes), width=draw_width, height=draw_height)
    doc.build([rl_image])


def export_full_page_image_docx(
    img: Image.Image, output_path: str, margin_cm: float = 0.8, page_size: str = "A4"
) -> None:
    """Embed `img` as a single, page-filling image in a new Word doc."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    width_cm, height_cm = _DOCX_PAGE_SIZES_CM[page_size]
    png_bytes = image_to_png_bytes(img)
    doc = Document()
    section = doc.sections[0]
    # Set explicitly rather than relying on Word's default template (which
    # isn't guaranteed to be A4/A3).
    section.page_width = Cm(width_cm)
    section.page_height = Cm(height_cm)
    section.left_margin = Cm(margin_cm)
    section.right_margin = Cm(margin_cm)
    section.top_margin = Cm(margin_cm)
    section.bottom_margin = Cm(margin_cm)
    usable_width_cm = (section.page_width - section.left_margin - section.right_margin) / Cm(1)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(io.BytesIO(png_bytes), width=Cm(usable_width_cm))
    doc.save(output_path)
