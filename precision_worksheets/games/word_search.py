"""Word Search using the target words: find them hidden in a grid of
letters, reading left-to-right, top-to-bottom, or diagonally down-right -
kept to those three "forwards" directions (no backwards or upwards words)
since this is aimed at early readers.
"""

from __future__ import annotations

import random
import string
from dataclasses import dataclass

DIRECTIONS = [(0, 1), (1, 0), (1, 1)]  # right, down, diagonal down-right
MIN_GRID_SIZE = 10
MAX_GRID_SIZE = 18
GRID_PADDING = 2  # extra rows/cols beyond the longest word, for placement variety


@dataclass
class WordSearchPuzzle:
    child_name: str
    words: list[str]  # original case, for the word list caption
    grid: list[list[str]]  # size x size of uppercase letters

    @property
    def size(self) -> int:
        return len(self.grid)


def build_word_search(child_name: str, words: list[str], seed: int | None = None) -> WordSearchPuzzle:
    clean_words = [w.strip() for w in words if w.strip()]
    if not clean_words:
        raise ValueError("At least one word is required")

    rng = random.Random(seed)
    longest = max(len(w) for w in clean_words)
    size = max(MIN_GRID_SIZE, min(MAX_GRID_SIZE, longest + GRID_PADDING))

    while True:
        try:
            grid = _try_build_grid(clean_words, size, rng)
            break
        except ValueError:
            if size >= MAX_GRID_SIZE:
                raise
            size += 2  # give placement more room and try again

    return WordSearchPuzzle(child_name=child_name, words=clean_words, grid=grid)


def _try_build_grid(words: list[str], size: int, rng: random.Random) -> list[list[str]]:
    grid: list[list[str | None]] = [[None] * size for _ in range(size)]
    for word in sorted(words, key=len, reverse=True):
        _place_word(grid, word.upper(), size, rng)

    alphabet = string.ascii_uppercase
    return [[cell if cell is not None else rng.choice(alphabet) for cell in row] for row in grid]


def _fits(grid: list[list[str | None]], word: str, row: int, col: int, dr: int, dc: int) -> bool:
    for i, letter in enumerate(word):
        existing = grid[row + i * dr][col + i * dc]
        if existing is not None and existing != letter:
            return False
    return True


def _place_word(grid: list[list[str | None]], word: str, size: int, rng: random.Random) -> None:
    positions = [
        (row, col, dr, dc)
        for dr, dc in DIRECTIONS
        for row in range(size - (len(word) - 1) * dr)
        for col in range(size - (len(word) - 1) * dc)
    ]
    rng.shuffle(positions)
    for row, col, dr, dc in positions:
        if _fits(grid, word, row, col, dr, dc):
            for i, letter in enumerate(word):
                grid[row + i * dr][col + i * dc] = letter
            return
    raise ValueError(f"Couldn't fit the word {word!r} on a {size}x{size} grid")


def export_word_search_pdf(puzzle: WordSearchPuzzle, output_path: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from ..fonts import get_comic_font_names
    from ..layout import fit_grid_font_size

    page_width, page_height = A4
    margin = 14 * mm
    doc = SimpleDocTemplate(
        output_path, pagesize=A4,
        topMargin=margin, bottomMargin=margin, leftMargin=margin, rightMargin=margin,
        title="Word Search",
    )
    comic_regular, comic_bold = get_comic_font_names()
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "WSTitle", parent=styles["Title"], fontName=comic_bold, fontSize=22,
        alignment=TA_CENTER, spaceAfter=4,
    )
    meta_style = ParagraphStyle(
        "WSMeta", parent=styles["Normal"], fontName=comic_regular, fontSize=11,
        alignment=TA_CENTER, spaceAfter=6,
    )
    instructions_style = ParagraphStyle(
        "WSInstructions", parent=styles["Normal"], fontName=comic_regular, fontSize=11,
        alignment=TA_CENTER, textColor=colors.HexColor("#444444"), spaceAfter=6,
    )
    word_list_style = ParagraphStyle(
        "WSWordList", parent=styles["Normal"], fontName=comic_bold, fontSize=13,
        alignment=TA_CENTER, spaceBefore=10,
    )

    usable_width = page_width - 2 * margin
    header_flowables = [
        Paragraph("Word Search", title_style),
        Paragraph(f"For {puzzle.child_name}", meta_style),
        Paragraph(
            "Find and circle every word below - they go across, down, "
            "or diagonally.", instructions_style,
        ),
        Spacer(1, 4),
    ]
    footer_flowable = Paragraph("Find these words: " + ", ".join(puzzle.words), word_list_style)
    header_height = sum(f.wrap(usable_width, page_height)[1] for f in header_flowables)
    footer_height = footer_flowable.wrap(usable_width, page_height)[1]

    cell_size = min(
        usable_width / puzzle.size,
        ((page_height - 2 * margin) - header_height - footer_height - 60) / puzzle.size,
    )
    letter_font_size = fit_grid_font_size(
        [letter for row in puzzle.grid for letter in row], int(cell_size * 0.6), cell_size,
    )

    data = puzzle.grid
    table = Table(data, colWidths=[cell_size] * puzzle.size, rowHeights=[cell_size] * puzzle.size)
    table.hAlign = "CENTER"
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), comic_bold),
                ("FONTSIZE", (0, 0), (-1, -1), letter_font_size),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )

    story = header_flowables + [table, footer_flowable]
    doc.build(story)


def export_word_search_docx(puzzle: WordSearchPuzzle, output_path: str) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm
    from reportlab.lib.units import cm as POINTS_PER_CM

    from ..docx_helpers import set_cell_text, set_column_widths, set_run_font
    from ..layout import fit_grid_font_size

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    usable_width_cm = (section.page_width - section.left_margin - section.right_margin) / Cm(1)

    title = doc.add_heading("Word Search", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        set_run_font(run)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(f"For {puzzle.child_name}")
    meta_run.bold = True
    set_run_font(meta_run)

    instructions = doc.add_paragraph()
    instructions.alignment = WD_ALIGN_PARAGRAPH.CENTER
    instructions_run = instructions.add_run(
        "Find and circle every word below - they go across, down, or diagonally."
    )
    instructions_run.italic = True
    set_run_font(instructions_run)

    table = doc.add_table(rows=puzzle.size, cols=puzzle.size)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell_width_cm = usable_width_cm / puzzle.size
    set_column_widths(table, [cell_width_cm] * puzzle.size)
    letter_font_size = fit_grid_font_size(
        [letter for row in puzzle.grid for letter in row], 24, cell_width_cm * POINTS_PER_CM,
    )

    for r, row_letters in enumerate(puzzle.grid):
        cells = table.rows[r].cells
        for c, letter in enumerate(row_letters):
            set_cell_text(cells[c], letter, size_pt=letter_font_size, bold=True)

    word_list = doc.add_paragraph()
    word_list.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word_list_run = word_list.add_run("Find these words: " + ", ".join(puzzle.words))
    word_list_run.bold = True
    set_run_font(word_list_run)

    doc.save(output_path)
