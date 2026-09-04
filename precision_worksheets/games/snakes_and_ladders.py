"""A printable Snakes & Ladders board using the target words.

Standard rules, plus one twist: a handful of squares show one of the
target words instead of a number - land on one and read it out loud; get
it right and take another turn. Everything (board, ladders, snakes, word
squares, instructions) is drawn once as a single image with Pillow, which
both the PDF and Word exporters just place full-page - so the two formats
always look identical.
"""

from __future__ import annotations

import io
import math
import random
from dataclasses import dataclass, field

from PIL import Image, ImageDraw, ImageFont

from ..fonts import get_board_font_paths

BOARD_SIZE = 10  # 10 x 10 = 100 squares
N_LADDERS = 6
N_SNAKES = 6
WORD_SQUARE_REPEATS = 2  # each word appears on this many board squares

# Layout, in "logical" pixels (drawn at SUPERSAMPLE x this and downscaled
# for anti-aliased edges on the diagonal ladders and curved snakes).
SUPERSAMPLE = 2
IMG_WIDTH = 1654
IMG_HEIGHT = 2339
CELL_SIZE = 150
BOARD_PIXELS = CELL_SIZE * BOARD_SIZE
BOARD_LEFT = (IMG_WIDTH - BOARD_PIXELS) // 2
BOARD_TOP = 220

COLOR_CELL_A = (255, 255, 255)
COLOR_CELL_B = (232, 244, 255)
COLOR_WORD_SQUARE = (255, 221, 130)
COLOR_GRID = (70, 70, 70)
COLOR_BORDER = (40, 40, 40)
COLOR_LADDER = (150, 100, 40)
COLOR_SNAKE = (60, 150, 95)
COLOR_TEXT = (30, 30, 30)
COLOR_MUTED = (90, 90, 90)


@dataclass
class SnakesAndLaddersBoard:
    child_name: str
    words: list[str]
    ladders: dict[int, int] = field(default_factory=dict)  # start -> end, climb up
    snakes: dict[int, int] = field(default_factory=dict)  # start -> end, slide down
    word_squares: dict[int, str] = field(default_factory=dict)  # square -> word


def build_board(child_name: str, words: list[str], seed: int | None = None) -> SnakesAndLaddersBoard:
    """Randomly place ladders, snakes and word squares on a 100-square board."""
    clean_words = [w.strip() for w in words if w.strip()]
    if not clean_words:
        raise ValueError("At least one word is required")

    rng = random.Random(seed)
    used: set[int] = {1, 100}
    ladders: dict[int, int] = {}
    snakes: dict[int, int] = {}

    def place(is_ladder: bool) -> None:
        for _ in range(200):
            if is_ladder:
                start = rng.randint(2, 88)
                end = min(99, start + rng.randint(10, 30))
            else:
                start = rng.randint(20, 98)
                end = max(2, start - rng.randint(10, 30))
            if start == end or start in used or end in used:
                continue
            if (end > start) != is_ladder:
                continue
            used.add(start)
            used.add(end)
            (ladders if is_ladder else snakes)[start] = end
            return

    for _ in range(N_LADDERS):
        place(True)
    for _ in range(N_SNAKES):
        place(False)

    n_word_squares = len(clean_words) * WORD_SQUARE_REPEATS
    word_pool = (clean_words * WORD_SQUARE_REPEATS)[:n_word_squares]
    rng.shuffle(word_pool)
    available = [n for n in range(2, 100) if n not in used]
    rng.shuffle(available)
    word_squares = dict(zip(available, word_pool))

    return SnakesAndLaddersBoard(
        child_name=child_name, words=clean_words,
        ladders=ladders, snakes=snakes, word_squares=word_squares,
    )


def _square_center(n: int) -> tuple[int, int]:
    """Pixel centre of square `n` (1-100), numbered boustrophedon-style:
    1-10 left to right along the bottom row, 11-20 right to left, and so on
    up to 100 - the classic Snakes & Ladders layout."""
    idx = n - 1
    row_from_bottom, col_in_row = divmod(idx, BOARD_SIZE)
    if row_from_bottom % 2 == 1:
        col_in_row = BOARD_SIZE - 1 - col_in_row
    row_from_top = BOARD_SIZE - 1 - row_from_bottom
    x = BOARD_LEFT + col_in_row * CELL_SIZE + CELL_SIZE // 2
    y = BOARD_TOP + row_from_top * CELL_SIZE + CELL_SIZE // 2
    return x, y


def _wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: float) -> list[str]:
    """Greedily wrap `text` onto lines that each fit within max_width."""
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        box = draw.textbbox((0, 0), candidate, font=font)
        if box[2] - box[0] <= max_width or not current:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def _fit_font(draw: ImageDraw.ImageDraw, text: str, font_path: str, max_size: int, max_width: int, min_size: int = 14) -> ImageFont.FreeTypeFont:
    size = max_size
    while size > min_size:
        font = ImageFont.truetype(font_path, size)
        box = draw.textbbox((0, 0), text, font=font)
        if box[2] - box[0] <= max_width:
            return font
        size -= 2
    return ImageFont.truetype(font_path, min_size)


def _perpendicular(dx: float, dy: float) -> tuple[float, float]:
    length = math.hypot(dx, dy) or 1
    return -dy / length, dx / length


def _draw_ladder(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    length = math.hypot(dx, dy)
    px, py = _perpendicular(dx, dy)
    rail_gap, width = 20, 10
    rail1 = [(x1 + px * rail_gap, y1 + py * rail_gap), (x2 + px * rail_gap, y2 + py * rail_gap)]
    rail2 = [(x1 - px * rail_gap, y1 - py * rail_gap), (x2 - px * rail_gap, y2 - py * rail_gap)]
    draw.line(rail1, fill=COLOR_LADDER, width=width)
    draw.line(rail2, fill=COLOR_LADDER, width=width)
    steps = max(2, int(length // 38))
    for i in range(steps + 1):
        t = i / steps
        cx, cy = x1 + dx * t, y1 + dy * t
        draw.line(
            [(cx + px * rail_gap, cy + py * rail_gap), (cx - px * rail_gap, cy - py * rail_gap)],
            fill=COLOR_LADDER, width=max(4, width - 3),
        )


def _draw_snake(draw: ImageDraw.ImageDraw, head: tuple[int, int], tail: tuple[int, int]) -> None:
    x1, y1 = head
    x2, y2 = tail
    dx, dy = x2 - x1, y2 - y1
    length = math.hypot(dx, dy)
    px, py = _perpendicular(dx, dy)
    waves = max(1, round(length / 130))
    n = max(20, int(length / 8))
    width = 15
    points = []
    for i in range(n + 1):
        t = i / n
        bx, by = x1 + dx * t, y1 + dy * t
        offset = math.sin(t * math.pi * waves) * 28 * (1 - 0.35 * t)
        points.append((bx + px * offset, by + py * offset))
    draw.line(points, fill=COLOR_SNAKE, width=width, joint="curve")
    hx, hy = points[0]
    r = width * 0.95
    draw.ellipse([hx - r, hy - r, hx + r, hy + r], fill=COLOR_SNAKE)
    eye_dx, eye_dy = px * r * 0.4, py * r * 0.4
    for sign in (-1, 1):
        ex, ey = hx + eye_dx * sign, hy + eye_dy * sign
        draw.ellipse([ex - 3, ey - 3, ex + 3, ey + 3], fill=(255, 255, 255))
    tx, ty = points[-1]
    draw.ellipse([tx - width * 0.3, ty - width * 0.3, tx + width * 0.3, ty + width * 0.3], fill=COLOR_SNAKE)


def render_board_image(board: SnakesAndLaddersBoard) -> Image.Image:
    """Draw the full board (title, grid, ladders, snakes, legend) as one image."""
    scale = SUPERSAMPLE
    img = Image.new("RGB", (IMG_WIDTH * scale, IMG_HEIGHT * scale), (255, 255, 255))
    draw = ImageDraw.Draw(img)

    def S(v: float) -> float:
        return v * scale

    regular_path, bold_path = get_board_font_paths()
    title_font = ImageFont.truetype(bold_path, int(S(56)))
    subtitle_font = ImageFont.truetype(regular_path, int(S(30)))
    number_font = ImageFont.truetype(regular_path, int(S(22)))
    legend_font = ImageFont.truetype(regular_path, int(S(26)))
    legend_title_font = ImageFont.truetype(bold_path, int(S(30)))

    def centered_text(y: float, text: str, font: ImageFont.FreeTypeFont, fill=COLOR_TEXT) -> None:
        box = draw.textbbox((0, 0), text, font=font)
        w = box[2] - box[0]
        draw.text((S(IMG_WIDTH / 2) - w / 2, S(y)), text, font=font, fill=fill)

    centered_text(50, "Snakes & Ladders", title_font)
    centered_text(120, f"for {board.child_name}", subtitle_font, fill=COLOR_MUTED)

    # Board squares
    for row in range(BOARD_SIZE):
        for col in range(BOARD_SIZE):
            n = _square_number_at(row, col)
            x0 = S(BOARD_LEFT + col * CELL_SIZE)
            y0 = S(BOARD_TOP + row * CELL_SIZE)
            x1 = x0 + S(CELL_SIZE)
            y1 = y0 + S(CELL_SIZE)
            is_word_square = n in board.word_squares
            if is_word_square:
                fill = COLOR_WORD_SQUARE
            else:
                fill = COLOR_CELL_A if (row + col) % 2 == 0 else COLOR_CELL_B
            draw.rectangle([x0, y0, x1, y1], fill=fill, outline=COLOR_GRID, width=max(1, int(S(1.5))))

            if is_word_square:
                word = board.word_squares[n]
                font = _fit_font(draw, word, bold_path, int(S(46)), int(S(CELL_SIZE - 24)))
                box = draw.textbbox((0, 0), word, font=font)
                tw, th = box[2] - box[0], box[3] - box[1]
                draw.text(
                    (x0 + S(CELL_SIZE) / 2 - tw / 2, y0 + S(CELL_SIZE) / 2 - th / 2 - box[1]),
                    word, font=font, fill=COLOR_TEXT,
                )
            else:
                draw.text((x0 + S(6), y0 + S(4)), str(n), font=number_font, fill=COLOR_MUTED)

    # Ladders and snakes, drawn over the grid.
    for start, end in board.ladders.items():
        _draw_ladder(draw, tuple(S(v) for v in _square_center(start)), tuple(S(v) for v in _square_center(end)))
    for start, end in board.snakes.items():
        _draw_snake(draw, tuple(S(v) for v in _square_center(start)), tuple(S(v) for v in _square_center(end)))

    # Outer board border, drawn last so it sits crisply on top.
    draw.rectangle(
        [S(BOARD_LEFT), S(BOARD_TOP), S(BOARD_LEFT + BOARD_PIXELS), S(BOARD_TOP + BOARD_PIXELS)],
        outline=COLOR_BORDER, width=max(2, int(S(3))),
    )

    # Legend / instructions below the board.
    legend_top = BOARD_TOP + BOARD_PIXELS + 50
    centered_text(legend_top, "How to play", legend_title_font)

    swatch = 26
    line_y = legend_top + 55
    draw.rectangle([S(BOARD_LEFT), S(line_y), S(BOARD_LEFT + swatch), S(line_y + swatch)], fill=COLOR_WORD_SQUARE, outline=COLOR_GRID)
    draw.text((S(BOARD_LEFT + swatch + 16), S(line_y)), "Land here? Read the word out loud - get it right and roll again!", font=legend_font, fill=COLOR_TEXT)

    line_y += 50
    draw.line([(S(BOARD_LEFT), S(line_y + swatch / 2)), (S(BOARD_LEFT + swatch), S(line_y + swatch / 2))], fill=COLOR_LADDER, width=int(S(8)))
    draw.text((S(BOARD_LEFT + swatch + 16), S(line_y)), "Climb a ladder", font=legend_font, fill=COLOR_TEXT)

    line_y += 50
    draw.line([(S(BOARD_LEFT), S(line_y + swatch / 2)), (S(BOARD_LEFT + swatch), S(line_y + swatch / 2))], fill=COLOR_SNAKE, width=int(S(8)))
    draw.text((S(BOARD_LEFT + swatch + 16), S(line_y)), "Slide down a snake", font=legend_font, fill=COLOR_TEXT)

    line_y += 60
    instructions = (
        "You'll need a die and a counter for each player. Take turns rolling "
        "and moving your counter that many squares. First to reach square "
        "100 wins!"
    )
    for line in _wrap_text(draw, instructions, legend_font, S(BOARD_PIXELS)):
        draw.text((S(BOARD_LEFT), S(line_y)), line, font=legend_font, fill=COLOR_MUTED)
        line_y += 38

    if scale != 1:
        img = img.resize((IMG_WIDTH, IMG_HEIGHT), Image.LANCZOS)
    return img


def _square_number_at(row: int, col: int) -> int:
    """Inverse of _square_center: the square number drawn at grid (row, col),
    where row 0 is the top row and col 0 is the left column."""
    row_from_bottom = BOARD_SIZE - 1 - row
    col_in_row = col if row_from_bottom % 2 == 0 else BOARD_SIZE - 1 - col
    return row_from_bottom * BOARD_SIZE + col_in_row + 1


def _image_to_png_bytes(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def export_board_pdf(board: SnakesAndLaddersBoard, output_path: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import Image as RLImage, SimpleDocTemplate

    img = render_board_image(board)
    png_bytes = _image_to_png_bytes(img)

    margin = 8 * mm
    doc = SimpleDocTemplate(
        output_path, pagesize=A4,
        topMargin=margin, bottomMargin=margin, leftMargin=margin, rightMargin=margin,
        title="Snakes and Ladders",
    )
    page_width, page_height = A4
    usable_width = page_width - 2 * margin
    usable_height = page_height - 2 * margin
    aspect = IMG_HEIGHT / IMG_WIDTH
    draw_width = usable_width
    draw_height = draw_width * aspect
    if draw_height > usable_height:
        draw_height = usable_height
        draw_width = draw_height / aspect

    rl_image = RLImage(io.BytesIO(png_bytes), width=draw_width, height=draw_height)
    doc.build([rl_image])


def export_board_docx(board: SnakesAndLaddersBoard, output_path: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    img = render_board_image(board)
    png_bytes = _image_to_png_bytes(img)

    doc = Document()
    section = doc.sections[0]
    # A4 portrait, set explicitly rather than relying on Word's default
    # template (which isn't guaranteed to be A4).
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(0.8)
    section.right_margin = Cm(0.8)
    section.top_margin = Cm(0.8)
    section.bottom_margin = Cm(0.8)
    usable_width_cm = (section.page_width - section.left_margin - section.right_margin) / Cm(1)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(io.BytesIO(png_bytes), width=Cm(usable_width_cm))
    doc.save(output_path)
