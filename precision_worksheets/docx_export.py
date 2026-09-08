"""Render ProbeSheet objects to a printable, editable Word document."""

from __future__ import annotations

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from reportlab.lib.units import cm as POINTS_PER_CM

from .docx_helpers import set_cell_margins, set_cell_text, set_column_widths, set_run_font
from .generator import ProbeSheet
from .layout import fit_grid_font_size

ROW_NUM_COL_CM = 0.8

# A4 portrait, set explicitly rather than relying on Word's default
# template (which isn't guaranteed to be A4).
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7

# Rough space the header (title/name/word list) takes up, used to work out
# how much vertical room is left for the grid itself - padded well over
# the true value so an estimation error never quietly pushes a sheet onto
# a near-blank extra page. Word's built-in "Heading 1" style carries its
# own template spacing that runs taller than its font size alone would
# suggest, so this needs more headroom than the equivalent estimate in
# pdf_export.py.
_HEADER_HEIGHT_WITH_WORDLIST_PT = 130
_HEADER_HEIGHT_WITHOUT_WORDLIST_PT = 105
_GRID_LEADING_FACTOR = 1.15
_MIN_ROW_PADDING_PT = 10

# The tracker table (when shown) gets roughly this fraction of the page's
# usable height, at the cost of a shorter grid above it - a fixed line
# without it just gets a little breathing room instead.
_TRACKER_HEIGHT_FRACTION = 0.25
_NO_TRACKER_FOOTER_HEIGHT_PT = 15
_TRACKER_TITLE_HEIGHT_PT = 40
_TRACKER_MIN_ROW_PADDING_PT = 6
# Slack subtracted off the grid's share of the page, so small estimation
# errors in either block's real rendered height can never push the
# tracker table into a split across two pages.
_SAFETY_MARGIN_PT = 25

TRACKER_TRIES = 10
_TRACKER_LABEL_COL_CM = 1.5


def export_docx(
    sheets: list[ProbeSheet],
    output_path: str,
    include_word_list: bool = True,
    grid_font_size: int = 20,
    include_tracker: bool = True,
) -> None:
    """Write `sheets` to `output_path` as a single multi-page .docx file.

    `grid_font_size` is a ceiling, not a fixed size: a word only shrinks
    below it if it wouldn't otherwise fit on one line in its box. The boxes
    themselves are padded to fill the space left on the page after the
    header/footer, so the grid uses as much of the sheet as it can without
    needing the text itself to be huge.

    `include_tracker` adds a compact progress-tracker table to the bottom
    of every sheet - one row to log the date and one to log the score
    across up to 10 tries at that sheet - at the cost of a shorter grid
    above it to make room.
    """
    if not sheets:
        raise ValueError("No sheets to export")

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)
    usable_width_emu = section.page_width - section.left_margin - section.right_margin
    usable_width_cm = usable_width_emu / Cm(1)

    for sheet in sheets:
        title = doc.add_heading("Precision Teaching Probe Sheet", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            set_run_font(run)

        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta.add_run(
            f"Name: {sheet.child_name}     Date: {sheet.date_str}     "
            f"Sheet {sheet.sheet_number} of {sheet.total_sheets}"
        )
        meta_run.bold = True
        set_run_font(meta_run)

        if include_word_list:
            wp = doc.add_paragraph()
            wp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = wp.add_run("Target words: " + ", ".join(sheet.words))
            run.italic = True
            set_run_font(run)

        table = doc.add_table(rows=sheet.rows, cols=sheet.cols + 1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.autofit = False

        col_width_cm = (usable_width_cm - ROW_NUM_COL_CM) / sheet.cols
        set_column_widths(table, [ROW_NUM_COL_CM] + [col_width_cm] * sheet.cols)

        # A word only shrinks below grid_font_size if it wouldn't otherwise
        # fit on one line in its box.
        fitted_font_size = fit_grid_font_size(
            sheet.words, grid_font_size, col_width_cm * POINTS_PER_CM
        )

        # Pad each row so the grid stretches to fill the space left on the
        # page below the header and (if shown) the tracker, rather than
        # leaving it mostly blank.
        header_height = (
            _HEADER_HEIGHT_WITH_WORDLIST_PT if include_word_list else _HEADER_HEIGHT_WITHOUT_WORDLIST_PT
        )
        page_height_pt = PAGE_HEIGHT_CM * POINTS_PER_CM
        margins_pt = (section.top_margin + section.bottom_margin) / Cm(1) * POINTS_PER_CM
        usable_height_pt = page_height_pt - margins_pt
        footer_height = (
            usable_height_pt * _TRACKER_HEIGHT_FRACTION if include_tracker else _NO_TRACKER_FOOTER_HEIGHT_PT
        )
        available_grid_height = usable_height_pt - header_height - footer_height - _SAFETY_MARGIN_PT
        target_row_height = available_grid_height / sheet.rows
        text_block_height = fitted_font_size * _GRID_LEADING_FACTOR
        vertical_padding = max(_MIN_ROW_PADDING_PT, (target_row_height - text_block_height) / 2)
        set_cell_margins(table, top_pt=vertical_padding, bottom_pt=vertical_padding, left_pt=6, right_pt=6)

        for r, row_words in enumerate(sheet.grid):
            cells = table.rows[r].cells
            set_cell_text(cells[0], str(r + 1), size_pt=9, bold=False)
            for c, word in enumerate(row_words, start=1):
                set_cell_text(cells[c], word, size_pt=fitted_font_size, bold=True)

        if include_tracker:
            tracker_title = doc.add_paragraph()
            tracker_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tracker_title_run = tracker_title.add_run("Progress Tracker")
            tracker_title_run.bold = True
            tracker_title_run.font.size = Pt(13)
            set_run_font(tracker_title_run)

            tracker_table = doc.add_table(rows=3, cols=TRACKER_TRIES + 1)
            tracker_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            tracker_table.style = "Table Grid"
            tracker_table.autofit = False
            try_col_width_cm = (usable_width_cm - _TRACKER_LABEL_COL_CM) / TRACKER_TRIES
            set_column_widths(tracker_table, [_TRACKER_LABEL_COL_CM] + [try_col_width_cm] * TRACKER_TRIES)

            tracker_row_area = footer_height - _TRACKER_TITLE_HEIGHT_PT
            tracker_target_row_height = tracker_row_area / 3
            tracker_text_height = 10 * _GRID_LEADING_FACTOR
            tracker_padding = max(
                _TRACKER_MIN_ROW_PADDING_PT, (tracker_target_row_height - tracker_text_height) / 2
            )
            set_cell_margins(tracker_table, top_pt=tracker_padding, bottom_pt=tracker_padding, left_pt=4, right_pt=4)

            try_row, date_row, score_row = tracker_table.rows
            set_cell_text(try_row.cells[0], "Try", size_pt=10, bold=True)
            for i in range(1, TRACKER_TRIES + 1):
                set_cell_text(try_row.cells[i], str(i), size_pt=10, bold=True)
            set_cell_text(date_row.cells[0], "Date", size_pt=10, bold=True)
            set_cell_text(score_row.cells[0], "Score", size_pt=10, bold=True)
            for i in range(1, TRACKER_TRIES + 1):
                set_cell_text(date_row.cells[i], "", size_pt=10, bold=False)
                set_cell_text(score_row.cells[i], "", size_pt=10, bold=False)

        if sheet.sheet_number != sheet.total_sheets:
            doc.add_page_break()

    doc.save(output_path)
