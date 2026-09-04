"""Small python-docx helpers shared by every Word exporter in this app.

python-docx doesn't expose cell padding or a way to lock column widths
through its normal API, so these reach into the underlying OOXML - kept
here once rather than duplicated in every exporter that builds a table.
"""

from __future__ import annotations

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# Comic Sans MS ships with every edition of Windows, so referencing it by
# name (rather than embedding a font file) is reliable for a Windows app -
# Word just uses whatever copy is already installed on the machine.
COMIC_FONT_NAME = "Comic Sans MS"


def set_run_font(run, name: str = COMIC_FONT_NAME) -> None:
    run.font.name = name
    # Word can pick a different font for the "complex script" slot even
    # when w:ascii/w:hAnsi are set, so pin that explicitly too.
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:cs"), name)


def set_cell_text(cell, text: str, size_pt: float, bold: bool) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    set_run_font(run)


def set_cell_margins(table, top_pt: float, bottom_pt: float, left_pt: float, right_pt: float) -> None:
    """python-docx has no direct API for cell padding, so set it via the
    table-wide tblCellMar - twentieths of a point (dxa) per the OOXML spec.

    tblPr's children must stay in schema order (tblLayout, then tblCellMar,
    then tblLook, ...), so this is inserted before tblLook rather than just
    appended - a plain append would land after tblLook, which python-docx's
    add_table() already added.
    """
    tbl_pr = table._tbl.tblPr
    cell_mar = OxmlElement("w:tblCellMar")
    for tag, value_pt in (("top", top_pt), ("bottom", bottom_pt), ("left", left_pt), ("right", right_pt)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(int(value_pt * 20)))
        node.set(qn("w:type"), "dxa")
        cell_mar.append(node)
    tbl_pr.insert_element_before(
        cell_mar, "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange"
    )


def set_column_widths(table, widths_cm: list[float]) -> None:
    """python-docx needs the width set on every cell, not just the column,
    for Word to respect it reliably. Setting autofit=False already adds a
    correctly-ordered `w:tblLayout type="fixed"` (see CT_TblPr.autofit), so
    there's no need to add a second one by hand."""
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
